from pathlib import Path
import shutil
import json
import pandas as pd
from openpyxl import load_workbook


# =============================================================================
# CONFIGURATION
# =============================================================================

ROOT = Path(__file__).resolve().parents[1]
META = ROOT / "data" / "metadata"

OUT_DIR = ROOT / "final_result_tables"
OUT_DIR.mkdir(parents=True, exist_ok=True)

SEEDS = list(range(42, 62))

TABLES = list(range(3, 18))

AUDIT_ROWS = []


# =============================================================================
# HELPERS
# =============================================================================

def banner(text):
    print("\n" + "=" * 110)
    print(text)
    print("=" * 110)


def normalize_text(x):
    if x is None:
        return ""
    return str(x).strip().lower()


def parse_horizon(x):
    if x is None:
        return None

    s = str(x).strip().lower()

    for token in [
        "hours", "hour", "hrs", "hr", "h"
    ]:
        s = s.replace(token, "")

    s = s.strip()

    try:
        return int(float(s))
    except Exception:
        return None


def record_change(
    table,
    sheet,
    row,
    column,
    seed,
    source,
    mapping,
    old_value,
    new_value,
):
    AUDIT_ROWS.append(
        {
            "table": table,
            "sheet": sheet,
            "row": row,
            "column": column,
            "seed": seed,
            "source_file": str(source),
            "mapping": mapping,
            "original_template_value": old_value,
            "real_replacement_value": new_value,
            "status": "REPLACED_WITH_REAL_RESULT",
        }
    )


def set_real_value(
    ws,
    row,
    col,
    value,
    *,
    table,
    seed,
    source,
    mapping,
):
    old = ws.cell(row, col).value

    ws.cell(row, col).value = float(value)

    record_change(
        table=table,
        sheet=ws.title,
        row=row,
        column=ws.cell(row, col).column_letter,
        seed=seed,
        source=source,
        mapping=mapping,
        old_value=old,
        new_value=float(value),
    )


# =============================================================================
# VERIFY INPUT TABLES
# =============================================================================

banner("STEP 34 - POPULATE ALL RESULT TABLES WITH AVAILABLE REAL RESULTS")

print(f"Project root:     {ROOT}")
print(f"Metadata root:    {META}")
print(f"Seeds:            {SEEDS[0]}-{SEEDS[-1]}")
print(f"Tables:           3-17")
print()
print("RULE:")
print("  Real result available + exact mapping -> REPLACE")
print("  Otherwise                            -> KEEP TEMPLATE VALUE")
print()


for table in TABLES:

    path = ROOT / f"Table{table}_20_Run.xlsx"

    if not path.exists():
        raise FileNotFoundError(
            f"Missing template:\n{path}"
        )

print("All 15 simple-name templates found.")


# =============================================================================
# LOAD REAL STEP-31 RESULTS
# =============================================================================

banner("LOADING REAL SEED 42-61 STEP-31 RESULTS")

overall_records = []
horizon_records = []

for seed in SEEDS:

    overall_path = (
        META
        / "target_test_analysis"
        / f"seed_{seed}"
        / "step31_common_subset_overall_metrics.csv"
    )

    horizon_path = (
        META
        / "target_test_analysis"
        / f"seed_{seed}"
        / "step31_common_subset_horizon_metrics.csv"
    )

    if not overall_path.exists():
        raise FileNotFoundError(overall_path)

    if not horizon_path.exists():
        raise FileNotFoundError(horizon_path)

    overall = pd.read_csv(overall_path)
    horizon = pd.read_csv(horizon_path)

    overall["seed"] = seed
    horizon["seed"] = seed

    overall["source_file"] = str(
        overall_path.relative_to(ROOT)
    )

    horizon["source_file"] = str(
        horizon_path.relative_to(ROOT)
    )

    overall_records.append(overall)
    horizon_records.append(horizon)


OVERALL = pd.concat(
    overall_records,
    ignore_index=True,
)

HORIZON = pd.concat(
    horizon_records,
    ignore_index=True,
)


print("Overall rows:", len(OVERALL))
print("Horizon rows:", len(HORIZON))

print(
    "Variants:",
    sorted(OVERALL["variant"].unique())
)


# =============================================================================
# LOAD OPTIONAL STEP-34A sMAPE RESULTS
# =============================================================================

SMAPE_PATH = (
    META
    / "table4_results"
    / "table4_existing_prediction_smape.csv"
)

if SMAPE_PATH.exists():

    SMAPE = pd.read_csv(SMAPE_PATH)

    print(
        "Step-34A sMAPE rows:",
        len(SMAPE),
    )

else:

    SMAPE = pd.DataFrame()

    print(
        "Step-34A sMAPE file not found."
    )

    print(
        "Table-4 sMAPE template values will therefore remain unchanged."
    )


# =============================================================================
# COPY ALL 15 ORIGINAL TEMPLATES FIRST
# =============================================================================

banner("COPYING TABLES 3-17")

OUTPUT_PATHS = {}

for table in TABLES:

    src = ROOT / f"Table{table}_20_Run.xlsx"

    dst = (
        OUT_DIR
        / f"Table{table}_20_Run_UPDATED.xlsx"
    )

    shutil.copy2(
        src,
        dst,
    )

    OUTPUT_PATHS[table] = dst

    print(
        f"Table {table:2d}: copied"
    )


# =============================================================================
# TABLE 3
# =============================================================================

banner("TABLE 3 - REAL PM2.5 HORIZON RMSE")

table = 3
path = OUTPUT_PATHS[table]

wb = load_workbook(path)

if "Results_Table3" not in wb.sheetnames:
    raise RuntimeError(
        "Results_Table3 sheet not found."
    )

ws = wb["Results_Table3"]


# Find headers.
headers = {
    normalize_text(
        ws.cell(1, col).value
    ): col
    for col in range(
        1,
        ws.max_column + 1,
    )
}


target_col = headers["target"]
method_col = headers["method"]
horizon_col = headers["horizon"]

run_cols = {}

for run in range(1, 21):

    key = f"run {run}"

    if key in headers:
        run_cols[run] = headers[key]
    elif f"seed {run}" in headers:
        run_cols[run] = headers[f"seed {run}"]
    else:
        raise RuntimeError(
            f"Table 3 missing Run/Seed {run} column."
        )


# Build row lookup.
row_lookup = {}

for row in range(
    2,
    ws.max_row + 1,
):

    target = normalize_text(
        ws.cell(row, target_col).value
    )

    method = normalize_text(
        ws.cell(row, method_col).value
    )

    horizon = parse_horizon(
        ws.cell(row, horizon_col).value
    )

    if not target or not method or horizon is None:
        continue

    row_lookup[
        (
            target,
            method,
            horizon,
        )
    ] = row


TABLE3_MAPPING = {
    "persistence": "M0",
    "proposed": "M2",
}


for method, variant in TABLE3_MAPPING.items():

    for horizon in [
        1,
        3,
        6,
        12,
        24,
    ]:

        # Try common PM2.5 labels.
        row = None

        for target_name in [
            "pm2.5",
            "pm25",
            "pm 2.5",
            "pm₂.₅",
        ]:

            key = (
                target_name,
                method,
                horizon,
            )

            if key in row_lookup:
                row = row_lookup[key]
                break

        if row is None:

            print(
                f"[Table 3] Row not found: "
                f"PM2.5 / {method} / {horizon}h"
            )

            continue


        for run_number, seed in enumerate(
            SEEDS,
            start=1,
        ):

            x = HORIZON[
                (HORIZON["seed"] == seed)
                &
                (HORIZON["variant"] == variant)
                &
                (
                    HORIZON["horizon_hours"]
                    .astype(int)
                    == horizon
                )
            ]

            if len(x) != 1:
                raise RuntimeError(
                    f"Table 3 lookup failure: "
                    f"seed={seed}, "
                    f"variant={variant}, "
                    f"horizon={horizon}"
                )

            value = float(
                x.iloc[0]["rmse_pm25"]
            )

            source = x.iloc[0]["source_file"]

            set_real_value(
                ws,
                row,
                run_cols[run_number],
                value,
                table="Table3",
                seed=seed,
                source=source,
                mapping=(
                    f"PM2.5 {method} "
                    f"{horizon}h RMSE <- {variant}"
                ),
            )


wb.save(path)

print(
    "Table 3 updated with available real results."
)


# =============================================================================
# TABLE 4
# =============================================================================

banner("TABLE 4 - REAL TRANSFER RESULTS")

table = 4
path = OUTPUT_PATHS[table]

wb = load_workbook(path)

if "Results_Table4" not in wb.sheetnames:
    raise RuntimeError(
        "Results_Table4 sheet not found."
    )

ws = wb["Results_Table4"]


headers = {
    normalize_text(
        ws.cell(1, col).value
    ): col
    for col in range(
        1,
        ws.max_column + 1,
    )
}


configuration_col = headers["configuration"]
metric_col = headers["metric"]

run_cols = {}

for run in range(1, 21):

    key = f"run {run}"

    if key not in headers:
        raise RuntimeError(
            f"Table 4 missing {key}."
        )

    run_cols[run] = headers[key]


row_lookup = {}

for row in range(
    2,
    ws.max_row + 1,
):

    configuration = normalize_text(
        ws.cell(
            row,
            configuration_col,
        ).value
    )

    metric = normalize_text(
        ws.cell(
            row,
            metric_col,
        ).value
    )

    if configuration and metric:

        row_lookup[
            (
                configuration,
                metric,
            )
        ] = row


TABLE4_CONFIGS = {
    "target-only": "M1",
    "adapted transfer": "M2",
}


TABLE4_METRICS = {
    "mae": "mae_pm25",
    "rmse": "rmse_pm25",
    "r2": "r2",
}


# -------------------------------------------------------------------------
# MAE / RMSE / R2
# -------------------------------------------------------------------------

for configuration, variant in TABLE4_CONFIGS.items():

    for metric_label, result_column in TABLE4_METRICS.items():

        key = (
            configuration,
            metric_label,
        )

        if key not in row_lookup:

            print(
                f"[Table 4] Row not found: "
                f"{configuration} / {metric_label}"
            )

            continue

        row = row_lookup[key]

        for run_number, seed in enumerate(
            SEEDS,
            start=1,
        ):

            x = OVERALL[
                (OVERALL["seed"] == seed)
                &
                (OVERALL["variant"] == variant)
            ]

            if len(x) != 1:
                raise RuntimeError(
                    f"Table 4 lookup failure: "
                    f"seed={seed}, "
                    f"variant={variant}"
                )

            value = float(
                x.iloc[0][result_column]
            )

            source = x.iloc[0]["source_file"]

            set_real_value(
                ws,
                row,
                run_cols[run_number],
                value,
                table="Table4",
                seed=seed,
                source=source,
                mapping=(
                    f"{configuration} "
                    f"{metric_label} <- "
                    f"{variant}/{result_column}"
                ),
            )


# -------------------------------------------------------------------------
# OPTIONAL M2 sMAPE FROM STEP 34A
# -------------------------------------------------------------------------

if len(SMAPE):

    key = (
        "adapted transfer",
        "smape (%)",
    )

    if key in row_lookup:

        row = row_lookup[key]

        x_all = SMAPE[
            SMAPE["variant"].astype(str)
            == "M2"
        ].copy()

        for run_number, seed in enumerate(
            SEEDS,
            start=1,
        ):

            x = x_all[
                x_all["seed"].astype(int)
                == seed
            ]

            if len(x) != 1:

                print(
                    f"[Table 4] "
                    f"M2 sMAPE unavailable "
                    f"for seed {seed}; "
                    f"template retained."
                )

                continue

            value = float(
                x.iloc[0]["value"]
            )

            source = (
                SMAPE_PATH.relative_to(ROOT)
            )

            set_real_value(
                ws,
                row,
                run_cols[run_number],
                value,
                table="Table4",
                seed=seed,
                source=source,
                mapping=(
                    "Adapted transfer "
                    "sMAPE <- frozen M2 predictions"
                ),
            )


wb.save(path)

print(
    "Table 4 updated with available real results."
)


# =============================================================================
# TABLES 5-17
# =============================================================================

banner("TABLES 5-17 - PRESERVING UNMAPPED TEMPLATE VALUES")

print(
    "No ambiguous M0-M5 -> template-label substitutions will be made."
)

print()

for table in range(
    5,
    18,
):

    print(
        f"Table {table:2d}: "
        f"original template values retained "
        f"unless/until an exact experiment mapping is established."
    )


# =============================================================================
# AUDIT EVERY OUTPUT WORKBOOK
# =============================================================================

banner("WRITING MASTER AUDIT")

audit = pd.DataFrame(
    AUDIT_ROWS
)

audit_path = (
    OUT_DIR
    / "step34_real_result_replacement_audit.csv"
)

audit.to_csv(
    audit_path,
    index=False,
)


summary = (
    audit
    .groupby(
        ["table", "mapping"],
        dropna=False,
    )
    .size()
    .reset_index(
        name="cells_replaced"
    )
)

summary_path = (
    OUT_DIR
    / "step34_real_result_replacement_summary.csv"
)

summary.to_csv(
    summary_path,
    index=False,
)


print()
print(summary.to_string(index=False))


# =============================================================================
# MANIFEST
# =============================================================================

manifest = {
    "step": 34,
    "purpose":
        "Populate available real results into Tables 3-17 while preserving "
        "all template values for unavailable/unmapped experiments.",
    "independent_seeds": SEEDS,
    "run_mapping": {
        f"Run {i + 1}": seed
        for i, seed in enumerate(SEEDS)
    },
    "replacement_rule":
        "Replace only when an exact real-result mapping exists; "
        "otherwise retain original template value.",
    "safe_mappings": {
        "Table3": {
            "PM2.5 Persistence horizon RMSE": "M0",
            "PM2.5 Proposed horizon RMSE": "M2",
        },
        "Table4": {
            "Target-only MAE/RMSE/R2": "M1",
            "Adapted transfer MAE/RMSE/R2": "M2",
            "Adapted transfer sMAPE":
                "Step-34A frozen M2 predictions if available",
        },
    },
    "explicitly_not_assumed": [
        "M1 is not LSTM",
        "M3 is not Transformer",
        "M4 is not Graph WaveNet",
        "M5 is not DCRNN",
        "M5 is not automatically 'without modality masking'",
        "M3 is not automatically 'without domain alignment'",
        "No source-only result is inferred from the source energy head",
        "No PICP is fabricated from point predictions",
        "No Shift Score is fabricated",
    ],
    "cells_replaced_with_real_results":
        int(len(audit)),
    "output_directory":
        str(OUT_DIR),
}


manifest_path = (
    OUT_DIR
    / "step34_population_manifest.json"
)

with open(
    manifest_path,
    "w",
    encoding="utf-8",
) as f:

    json.dump(
        manifest,
        f,
        indent=2,
    )


# =============================================================================
# FINAL VERIFICATION
# =============================================================================

banner("STEP 34 COMPLETE")

print(
    f"Real result cells replaced: "
    f"{len(audit):,}"
)

print()
print("Generated workbooks:")

for table in TABLES:

    print(
        f"  {OUTPUT_PATHS[table]}"
    )

print()
print("Audit:")
print(
    f"  {audit_path}"
)

print(
    f"  {summary_path}"
)

print(
    f"  {manifest_path}"
)